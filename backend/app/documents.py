import io
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pymupdf as fitz
import pytesseract
from docx import Document
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from .config import settings

if settings.tesseract_cmd:
    pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd


def ocr(image):
    try:
        if image.width * image.height > 30_000_000:
            raise ValueError('An image exceeds the 30 megapixel OCR limit.')
        raw = pytesseract.image_to_string(image.convert('RGB'), lang='eng', timeout=25).strip()
        return '\n\n'.join(' '.join(p.splitlines()) for p in raw.split('\n\n'))
    except pytesseract.TesseractNotFoundError as exc:
        raise ValueError('OCR requires Tesseract. Use the Docker deployment or install Tesseract locally.') from exc
    except RuntimeError as exc:
        raise ValueError('OCR timed out. Try a smaller document or clearer scan.') from exc


def extract_document(filename, data):
    ext = Path(filename).suffix.lower()
    parts, ocr_count, page_count = [], 0, 0
    try:
        if ext == '.pdf':
            if not data.startswith(b'%PDF-'):
                raise ValueError('The file is not a valid PDF.')
            with fitz.open(stream=data, filetype='pdf') as doc:
                if doc.needs_pass:
                    raise ValueError('Password-protected PDFs are not supported. Upload an unlocked copy.')
                if len(doc) > settings.max_pages:
                    raise ValueError(f'Upload at most {settings.max_pages} PDF pages at a time.')
                page_count = len(doc)
                for page in doc:
                    # PDF line wrapping is visual, not a sentence boundary.
                    text = '\n'.join(' '.join(block[4].splitlines()) for block in page.get_text('blocks')
                                     if block[6] == 0).strip()
                    if len(re.sub(r'\W', '', text)) < 40:
                        scale = min(2, 2400 / max(page.rect.width, page.rect.height))
                        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                        text = ocr(Image.open(io.BytesIO(pix.tobytes('png'))))
                        ocr_count += 1
                    else:
                        # Mixed PDFs: OCR substantial embedded images as well as native text.
                        for image_info in page.get_images(full=True)[:10]:
                            image = doc.extract_image(image_info[0])
                            with Image.open(io.BytesIO(image['image'])) as im:
                                if im.width >= 300 and im.height >= 100:
                                    text += '\n' + ocr(im)
                                    ocr_count += 1
                    parts.append(text)
        elif ext == '.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('The expanded Word document exceeds 50 MB.')
                media = [n for n in archive.namelist() if n.startswith('word/media/')]
                if len(media) > 30:
                    raise ValueError('Word documents may contain at most 30 embedded images.')
                doc = Document(io.BytesIO(data))
                parts.extend(p.text for p in doc.paragraphs)
                for table in doc.tables:
                    parts.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
                for section in doc.sections:
                    parts.extend(p.text for p in section.header.paragraphs)
                    parts.extend(p.text for p in section.footer.paragraphs)
                for name in media:
                    if Path(name).suffix.lower() in {'.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.tif'}:
                        with Image.open(io.BytesIO(archive.read(name))) as im:
                            parts.append(ocr(im))
                            ocr_count += 1
        else:
            raise ValueError('Upload a PDF or .docx Word document. Save legacy .doc files as .docx first.')
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('The document could not be read. Check that it is a valid, unencrypted PDF or DOCX.') from exc
    text = '\n'.join(p.strip() for p in parts if p.strip())
    if not text:
        raise ValueError('No readable text was found. Try a clearer scan with English text.')
    if len(text) > settings.max_text_chars:
        raise ValueError(f'Extracted text exceeds {settings.max_text_chars:,} characters. Split the document.')
    return text, {'ocr_images_or_pages': ocr_count, 'pages': page_count, 'characters': len(text)}


def export_pdf(title, questions, answers):
    buffer = io.BytesIO()
    font = 'Helvetica'
    for path in ['/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 'C:/Windows/Fonts/arial.ttf']:
        if Path(path).exists():
            if 'SourceSans' not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont('SourceSans', path))
            font = 'SourceSans'
            break
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
        if hasattr(style, 'leading'):
            style.leading = max(style.leading, 15)
    story = [Paragraph(escape(title), styles['Title']),
             Paragraph('Question set' + (' with answers' if answers else ''), styles['Normal']), Spacer(1, 20)]
    for i, q in enumerate(questions, 1):
        story.append(Paragraph(f'{i}. [{q.kind.upper()}] {escape(q.prompt)}', styles['Heading3']))
        for j, option in enumerate(q.options):
            story.append(Paragraph(f'{chr(65+j)}. {escape(option)}', styles['Normal']))
        if answers:
            story.append(Paragraph('Answer: ' + escape(q.answer).replace('\n', '<br/>'), styles['Normal']))
        story.append(Spacer(1, 16))

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor('#64748b'))
        canvas.drawString(42, 25, 'Question Garden · Review generated material before assessment use')
        canvas.drawRightString(550, 25, str(doc.page))

    SimpleDocTemplate(buffer, rightMargin=42, leftMargin=42, topMargin=40, bottomMargin=42).build(
        story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
